import io
import os
import csv
import json
import zipfile
import logging
from typing import Dict, Any, Tuple, Optional
from PIL import Image

logger = logging.getLogger("veil.services.extraction")

def extract_and_preview_file(filename: str, file_bytes: bytes) -> Dict[str, Any]:
    """
    Analyzes an uploaded file, extracts searchable text, and creates preview data.
    Returns:
    {
        "extracted_text": str,
        "preview_type": "pdf" | "image" | "table" | "json" | "code" | "text" | "archive" | "unsupported",
        "preview_data": dict or None,
        "mime_type": str,
        "char_count": int,
        "line_count": int,
        "is_searchable": bool
    }
    """
    ext = os.path.splitext(filename)[1].lower()
    
    # 1. PDF
    if ext == ".pdf":
        return extract_pdf(file_bytes)

    # 2. DOCX
    elif ext in [".docx", ".doc"]:
        return extract_docx(file_bytes, ext)

    # 3. Spreadsheets (CSV, XLSX, XLS)
    elif ext in [".csv", ".tsv"]:
        return extract_csv(file_bytes, ext)
    elif ext in [".xlsx", ".xls"]:
        return extract_xlsx(file_bytes)

    # 4. Structured Data (JSON, XML)
    elif ext == ".json":
        return extract_json(file_bytes)
    elif ext == ".xml":
        return extract_xml(file_bytes)

    # 5. Markdown & Plain Text
    elif ext in [".md", ".markdown", ".mdown"]:
        text = decode_text(file_bytes)
        return {
            "extracted_text": text,
            "preview_type": "text",
            "preview_data": {"format": "markdown"},
            "mime_type": "text/markdown",
            "char_count": len(text),
            "line_count": len(text.splitlines()),
            "is_searchable": True
        }
    elif ext in [".txt", ".rtf", ".log"]:
        text = decode_text(file_bytes)
        return {
            "extracted_text": text,
            "preview_type": "text",
            "preview_data": {"format": "plain"},
            "mime_type": "text/plain",
            "char_count": len(text),
            "line_count": len(text.splitlines()),
            "is_searchable": True
        }

    # 6. Code & Configuration Files
    elif ext in [".py", ".ts", ".tsx", ".js", ".jsx", ".java", ".c", ".cpp", ".cs", ".go", ".rs", ".php", ".rb", ".sql", ".html", ".css", ".yaml", ".yml", ".toml", ".sh"]:
        text = decode_text(file_bytes)
        lang = ext.lstrip(".")
        return {
            "extracted_text": text,
            "preview_type": "code",
            "preview_data": {"language": lang, "line_count": len(text.splitlines())},
            "mime_type": f"text/x-{lang}",
            "char_count": len(text),
            "line_count": len(text.splitlines()),
            "is_searchable": True
        }

    # 7. Images
    elif ext in [".png", ".jpg", ".jpeg", ".webp", ".gif", ".svg"]:
        return extract_image(file_bytes, ext)

    # 8. Archives (ZIP)
    elif ext == ".zip":
        return extract_zip(file_bytes)

    # 9. Fallback attempt for arbitrary text
    else:
        try:
            text = decode_text(file_bytes)
            # If text has high ratio of printable characters
            printable_ratio = sum(c.isprintable() or c.isspace() for c in text[:1000]) / max(len(text[:1000]), 1)
            if printable_ratio > 0.90:
                return {
                    "extracted_text": text,
                    "preview_type": "text",
                    "preview_data": {"format": "general_text"},
                    "mime_type": "text/plain",
                    "char_count": len(text),
                    "line_count": len(text.splitlines()),
                    "is_searchable": True
                }
        except Exception:
            pass

        # Unsupported Binary File
        return {
            "extracted_text": "",
            "preview_type": "unsupported",
            "preview_data": {"reason": "Binary format preview unavailable"},
            "mime_type": "application/octet-stream",
            "char_count": 0,
            "line_count": 0,
            "is_searchable": False
        }

def decode_text(file_bytes: bytes) -> str:
    for enc in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
        try:
            return file_bytes.decode(enc).replace("\r\n", "\n").replace("\r", "\n")
        except UnicodeDecodeError:
            continue
    raise ValueError("Unable to decode text with supported character encodings.")

def extract_pdf(file_bytes: bytes) -> Dict[str, Any]:
    """
    Extracts text page-by-page from an uploaded PDF.
    Preserves 1-indexed page numbers.
    Detects scanned PDFs where no extractable text exists.
    """
    import pymupdf
    try:
        doc = pymupdf.open(stream=file_bytes, filetype="pdf")
        num_pages = len(doc)
        pages_data = []
        full_text_list = []
        total_chars = 0

        for page_idx in range(num_pages):
            page = doc[page_idx]
            raw_text = page.get_text("text") or ""
            # Clean text: remove null bytes, normalize whitespace
            raw_text = raw_text.replace("\x00", "")
            cleaned = "\n".join([line.rstrip() for line in raw_text.splitlines() if line.strip()])
            if cleaned:
                pages_data.append({
                    "page": page_idx + 1,
                    "text": cleaned
                })
                full_text_list.append(f"--- Page {page_idx + 1} ---\n{cleaned}")
                total_chars += len(cleaned)

        doc.close()

        # PART 7: Detect scanned PDF
        is_scanned = (num_pages > 0 and len(pages_data) == 0) or (num_pages > 0 and total_chars < 20)
        if is_scanned:
            logger.warning("Scanned PDF detected — text extraction requires OCR.")
            return {
                "extracted_text": "",
                "pages": [],
                "preview_type": "pdf",
                "preview_data": {
                    "page_count": num_pages,
                    "is_scanned": True,
                    "error": "Scanned PDF detected — text extraction requires OCR."
                },
                "mime_type": "application/pdf",
                "char_count": 0,
                "line_count": 0,
                "is_searchable": False,
                "is_scanned": True,
                "processing_status": "Scanned PDF detected — text extraction requires OCR."
            }

        full_text = "\n\n".join(full_text_list)
        return {
            "extracted_text": full_text,
            "pages": pages_data,
            "preview_type": "pdf",
            "preview_data": {
                "page_count": num_pages,
                "pages": [{"page": p["page"], "char_count": len(p["text"])} for p in pages_data],
                "is_scanned": False
            },
            "mime_type": "application/pdf",
            "char_count": total_chars,
            "line_count": len(full_text.splitlines()),
            "is_searchable": True,
            "is_scanned": False,
            "processing_status": "ready"
        }
    except Exception as e:
        logger.error(f"PyMuPDF extraction error, attempting pypdf fallback: {e}")
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        num_pages = len(reader.pages)
        pages_data = []
        full_text_list = []
        total_chars = 0
        for idx, page in enumerate(reader.pages):
            raw_text = page.extract_text() or ""
            raw_text = raw_text.replace("\x00", "")
            cleaned = "\n".join([line.rstrip() for line in raw_text.splitlines() if line.strip()])
            if cleaned:
                pages_data.append({"page": idx + 1, "text": cleaned})
                full_text_list.append(f"--- Page {idx + 1} ---\n{cleaned}")
                total_chars += len(cleaned)

        is_scanned = (num_pages > 0 and len(pages_data) == 0) or (num_pages > 0 and total_chars < 20)
        if is_scanned:
            return {
                "extracted_text": "",
                "pages": [],
                "preview_type": "pdf",
                "preview_data": {
                    "page_count": num_pages,
                    "is_scanned": True,
                    "error": "Scanned PDF detected — text extraction requires OCR."
                },
                "mime_type": "application/pdf",
                "char_count": 0,
                "line_count": 0,
                "is_searchable": False,
                "is_scanned": True,
                "processing_status": "Scanned PDF detected — text extraction requires OCR."
            }

        full_text = "\n\n".join(full_text_list)
        return {
            "extracted_text": full_text,
            "pages": pages_data,
            "preview_type": "pdf",
            "preview_data": {
                "page_count": num_pages,
                "pages": [{"page": p["page"], "char_count": len(p["text"])} for p in pages_data],
                "is_scanned": False
            },
            "mime_type": "application/pdf",
            "char_count": total_chars,
            "line_count": len(full_text.splitlines()),
            "is_searchable": True,
            "is_scanned": False,
            "processing_status": "ready"
        }

def extract_docx(file_bytes: bytes, ext: str) -> Dict[str, Any]:
    if ext == ".docx":
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        full_text = "\n\n".join(paragraphs)
        return {
            "extracted_text": full_text,
            "preview_type": "text",
            "preview_data": {"format": "docx", "paragraphs_count": len(paragraphs)},
            "mime_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "char_count": len(full_text),
            "line_count": len(paragraphs),
            "is_searchable": bool(full_text.strip())
        }
    else:
        # Legacy .doc fallback
        text = decode_text(file_bytes) if len(file_bytes) < 500000 else ""
        return {
            "extracted_text": text,
            "preview_type": "text",
            "preview_data": {"format": "doc"},
            "mime_type": "application/msword",
            "char_count": len(text),
            "line_count": len(text.splitlines()),
            "is_searchable": bool(text.strip())
        }

def extract_csv(file_bytes: bytes, ext: str) -> Dict[str, Any]:
    text = decode_text(file_bytes)
    delimiter = "\t" if ext == ".tsv" else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = []
    for i, r in enumerate(reader):
        if i > 50:  # Preview cap
            break
        rows.append(r)
    
    headers = rows[0] if rows else []
    data_rows = rows[1:30] if len(rows) > 1 else []

    # Format into text representation for Moss
    text_repr = "\n".join([", ".join(r) for r in rows[:100]])
    return {
        "extracted_text": text_repr,
        "preview_type": "table",
        "preview_data": {
            "headers": headers,
            "columns": headers,
            "rows": data_rows,
            "total_preview_rows": len(rows)
        },
        "mime_type": "text/csv",
        "char_count": len(text),
        "line_count": len(rows),
        "is_searchable": True
    }

def extract_xlsx(file_bytes: bytes) -> Dict[str, Any]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    sheet = wb.active
    rows = []
    for r in sheet.iter_rows(values_only=True):
        if any(cell is not None for cell in r):
            rows.append([str(c) if c is not None else "" for c in r])
        if len(rows) > 50:
            break

    headers = rows[0] if rows else []
    data_rows = rows[1:30] if len(rows) > 1 else []
    text_repr = "\n".join([", ".join(r) for r in rows[:100]])

    return {
        "extracted_text": text_repr,
        "preview_type": "table",
        "preview_data": {
            "headers": headers,
            "columns": headers,
            "rows": data_rows,
            "sheet_name": sheet.title if sheet else "Sheet1",
            "total_preview_rows": len(rows)
        },
        "mime_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "char_count": len(text_repr),
        "line_count": len(rows),
        "is_searchable": True
    }

def extract_json(file_bytes: bytes) -> Dict[str, Any]:
    text = decode_text(file_bytes)
    try:
        parsed = json.loads(text)
        formatted = json.dumps(parsed, indent=2)
        return {
            "extracted_text": formatted,
            "preview_type": "json",
            "preview_data": {"is_valid": True, "top_keys": list(parsed.keys()) if isinstance(parsed, dict) else []},
            "mime_type": "application/json",
            "char_count": len(formatted),
            "line_count": len(formatted.splitlines()),
            "is_searchable": True
        }
    except Exception:
        return {
            "extracted_text": text,
            "preview_type": "text",
            "preview_data": {"format": "json_raw"},
            "mime_type": "application/json",
            "char_count": len(text),
            "line_count": len(text.splitlines()),
            "is_searchable": True
        }

def extract_xml(file_bytes: bytes) -> Dict[str, Any]:
    text = decode_text(file_bytes)
    return {
        "extracted_text": text,
        "preview_type": "code",
        "preview_data": {"language": "xml", "line_count": len(text.splitlines())},
        "mime_type": "application/xml",
        "char_count": len(text),
        "line_count": len(text.splitlines()),
        "is_searchable": True
    }

def extract_image(file_bytes: bytes, ext: str) -> Dict[str, Any]:
    try:
        img = Image.open(io.BytesIO(file_bytes))
        width, height = img.size
        img_format = img.format or ext.lstrip(".").upper()
        return {
            "extracted_text": f"[Image file: {ext.lstrip('.').upper()}, Dimensions: {width}x{height}]",
            "preview_type": "image",
            "preview_data": {"width": width, "height": height, "format": img_format},
            "mime_type": f"image/{ext.lstrip('.')}",
            "char_count": 0,
            "line_count": 0,
            "is_searchable": False
        }
    except Exception:
        return {
            "extracted_text": "",
            "preview_type": "image",
            "preview_data": {},
            "mime_type": f"image/{ext.lstrip('.')}",
            "char_count": 0,
            "line_count": 0,
            "is_searchable": False
        }

def extract_zip(file_bytes: bytes) -> Dict[str, Any]:
    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes), "r") as z:
            file_list = z.namelist()
            text_repr = f"Archive contents:\n" + "\n".join(file_list[:100])
            return {
                "extracted_text": text_repr,
                "preview_type": "archive",
                "preview_data": {"files": file_list[:50], "total_files": len(file_list)},
                "mime_type": "application/zip",
                "char_count": len(text_repr),
                "line_count": len(file_list),
                "is_searchable": True
            }
    except Exception as e:
        return {
            "extracted_text": "",
            "preview_type": "unsupported",
            "preview_data": {"error": str(e)},
            "mime_type": "application/zip",
            "char_count": 0,
            "line_count": 0,
            "is_searchable": False
        }
